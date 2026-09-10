"""Проверка pipeline данных с двух сторон.

Зелёная проверка ничего не доказывает, пока не падала на заведомо ломаном
случае, поэтому каждый инвариант `check_data.py` здесь ломают намеренно и
убеждаются, что он срабатывает — и что на корректной выгрузке он молчит.

Проверки берутся из реестра `check_data`, а не перечисляются заново: иначе
инвариант, потерянный при рефакторинге, оставил бы этот прогон зелёным.

Запуск: `python3 tools/test_pipeline.py`.
"""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from dataclasses import replace
from datetime import date
from pathlib import Path

import check_data as cd
import voc_schema as vs

TOOLS = Path(__file__).resolve().parent
ROOT = TOOLS.parent


class Run:
    def __init__(self) -> None:
        self.failed = 0

    def expect(self, condition: bool, description: str, detail: str = "") -> None:
        if condition:
            print(f"ОК      {description}")
        else:
            self.failed += 1
            print(f"ПРОВАЛ  {description}")
            if detail:
                print(f"        {detail}")


def run_tool(name: str, *arguments: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, str(TOOLS / name), *arguments], capture_output=True, text=True
    )


def copy_of(extract: vs.Extract) -> vs.Extract:
    return vs.Extract(
        ratings=list(extract.ratings),
        dictionaries={channel: dict(mapping) for channel, mapping in extract.dictionaries.items()},
        plan=list(extract.plan),
        labels=list(extract.labels),
    )


def marks_summing(count: int, total: int) -> list[int]:
    base, rest = divmod(total, count)
    return [base + 1] * rest + [base] * (count - rest)


def mini_extract(slices: dict[tuple[str, str, str], tuple[int, int]]) -> vs.Extract:
    """Крошечная выгрузка с заданными объёмами и суммами оценок.

    Нужна там, где нарушение — свойство арифметики, а не строки: подобрать
    его мутацией 28 тысяч строк дороже, чем задать двумя числами.
    """
    ratings = []
    for (channel, period, segment), (count, total) in slices.items():
        year, month = (int(part) for part in period.split("-"))
        for index, mark in enumerate(marks_summing(count, total)):
            ratings.append(
                vs.Rating(
                    sheet=channel,
                    row=index + 2,
                    voc_ccode=f"{channel}-{period}-{segment}-{index}",
                    appeal_date=date(year, month, 1 + index % 28),
                    answer_date=None,
                    channel=channel,
                    segment=segment,
                    operation_name="Отправка платежа",
                    mark=mark,
                    client_comment=None,
                    oslk_expertise_name=None,
                    domain=None,
                    problem_type=None,
                    problem=None,
                )
            )
    channels = sorted({rating.channel for rating in ratings})
    return vs.Extract(
        ratings=ratings,
        dictionaries={channel: {"Отправка платежа": ("КП", "Платежи")} for channel in channels},
        plan=[],
        labels=[],
    )


ANCHOR = (cd.ANCHOR_CHANNEL, cd.ANCHOR_PERIOD)


def unregistered_checks() -> list[str]:
    """Проверки, написанные, но не попавшие в реестр.

    Реестр — единственный источник и для программы, и для этих тестов, но
    сам факт «функция написана» он не гарантирует: без этой сверки новый
    инвариант мог бы никогда не запускаться.
    """
    registered = set(cd.ALL_CHECKS.values())
    return sorted(
        name
        for name, value in vars(cd).items()
        if name.startswith("check_") and callable(value) and value not in registered
    )


def anchor_indexes(extract: vs.Extract) -> list[int]:
    return [
        index
        for index, rating in enumerate(extract.ratings)
        if (rating.channel, cd.period_of(rating)) == ANCHOR
    ]


# --- поломки свойств любой выгрузки -----------------------------------------


def break_channel_sheet(extract):
    extract.ratings[0] = replace(extract.ratings[0], channel="ДРУГОЙ")


def break_required_fields(extract):
    extract.ratings[5] = replace(extract.ratings[5], segment="")


def break_ccode_unique(extract):
    extract.ratings[1] = replace(extract.ratings[1], voc_ccode=extract.ratings[0].voc_ccode)


def break_mark_range(extract):
    extract.ratings[3] = replace(extract.ratings[3], mark=0)


def break_distribution(extract):
    extract.ratings[7] = replace(extract.ratings[7], mark=6)


def break_unmapped_trigger(extract):
    channel = next(iter(extract.dictionaries))
    del extract.dictionaries[channel][next(iter(extract.dictionaries[channel]))]


def break_plan_rows(extract):
    extract.plan[0] = replace(extract.plan[0], minimum=4.95, maximum=4.85)


def break_labels(extract):
    extract.labels.append(replace(extract.labels[0], row=99))


# --- поломки обещаний сэмпла ------------------------------------------------


def break_voc_anchor_mean(extract):
    moved = 0
    for index in anchor_indexes(extract):
        if extract.ratings[index].mark == 5 and moved < 300:
            extract.ratings[index] = replace(extract.ratings[index], mark=3)
            moved += 1


def break_voc_anchor_count(extract):
    for index in reversed(anchor_indexes(extract)[-5:]):
        del extract.ratings[index]


def break_period_coverage(extract):
    hole = date(2026, 5, 15)
    extract.ratings = [rating for rating in extract.ratings if rating.appeal_date != hole]


def break_delta_signs(extract):
    """Остаётся один канал — падения VOC нет ни у кого."""
    extract.ratings = [r for r in extract.ratings if r.channel == cd.ANCHOR_CHANNEL]


def break_delta_signs_segments(extract):
    """Все сегменты двигаются одинаково — вид заглушки макета."""
    extract.ratings = [replace(rating, mark=5) for rating in extract.ratings]


def break_plan_states(extract):
    extract.plan = [replace(row, minimum=1.0, maximum=5.0) for row in extract.plan]


def break_corridor_crossing(extract):
    extract.plan = [
        replace(row, minimum=1.0, maximum=5.0) if row.segment is None else row
        for row in extract.plan
    ]


def break_dip_day(extract):
    extract.ratings = [
        replace(rating, mark=5)
        if rating.appeal_date and rating.appeal_date.isoformat() == cd.DIP_DAY
        else rating
        for rating in extract.ratings
    ]


def break_antidriver(extract):
    """Продукт, который просел, поднимается обратно — рейтинг пуст."""
    extract.ratings = [
        replace(rating, mark=5)
        if cd.period_of(rating) == cd.ANCHOR_PERIOD and "Выписки" in cd.products_of(extract, rating)
        else rating
        for rating in extract.ratings
    ]


def break_empty_slice(extract):
    extract.ratings = [rating for rating in extract.ratings if rating.segment != "КИБ"]


def break_small_sample(extract):
    extract.ratings = [
        replace(rating, mark=5) if rating.segment == "КИБ" else rating
        for rating in extract.ratings
    ]


def break_dictionaries_differ(extract):
    left, right = sorted(extract.dictionaries)[:2]
    for trigger, value in extract.dictionaries[left].items():
        if trigger in extract.dictionaries[right]:
            extract.dictionaries[right][trigger] = value


def break_multi_operation(extract):
    extract.ratings = [
        replace(rating, operation_name=rating.operations[0]) for rating in extract.ratings
    ]


def break_blank_markup(extract):
    extract.ratings = [
        replace(rating, domain="Платежи", problem_type="Ошибка", problem="Платёж висит")
        for rating in extract.ratings
    ]


def break_label_gaps(extract):
    extract.labels = extract.labels + [
        vs.LabelRow(row=100 + index, dimension="problem_type", code=code, label=code)
        for index, code in enumerate(
            {rating.problem_type for rating in extract.ratings if rating.problem_type}
        )
    ]


def break_verbatim(extract):
    extract.ratings = [
        replace(rating, client_comment="Коротко.") if rating.client_comment else rating
        for rating in extract.ratings
    ]


CASES = [
    ("channel-sheet", break_channel_sheet),
    ("required-fields", break_required_fields),
    ("ccode-unique", break_ccode_unique),
    ("mark-range", break_mark_range),
    ("distribution", break_distribution),
    ("unmapped-trigger", break_unmapped_trigger),
    ("plan-rows", break_plan_rows),
    ("labels", break_labels),
    ("voc-anchor", break_voc_anchor_mean),
    ("voc-anchor", break_voc_anchor_count),
    ("period-coverage", break_period_coverage),
    ("delta-signs", break_delta_signs),
    ("delta-signs", break_delta_signs_segments),
    ("plan-states", break_plan_states),
    ("corridor-crossing", break_corridor_crossing),
    ("dip-day", break_dip_day),
    ("antidriver", break_antidriver),
    ("empty-slice", break_empty_slice),
    ("small-sample", break_small_sample),
    ("dictionary-differs", break_dictionaries_differ),
    ("multi-operation", break_multi_operation),
    ("blank-markup", break_blank_markup),
    ("label-gaps", break_label_gaps),
    ("verbatim", break_verbatim),
]


def check_arithmetic_cases(run: Run) -> None:
    """Инварианты, нарушение которых — свойство чисел, а не строки."""
    inconsistent = mini_extract(
        {
            (cd.ANCHOR_CHANNEL, cd.ANCHOR_PERIOD, "ММБ"): (8, 37),
            (cd.ANCHOR_CHANNEL, cd.PREVIOUS_PERIOD, "ММБ"): (8, 38),
        }
    )
    run.expect(
        bool(cd.check_delta_consistency(inconsistent)),
        "delta-consistency падает, когда дельта не совпадает с показанными значениями",
        f"4.625 → 4.62 и 4.75: дельта −0.12, а по показанным −0.13",
    )
    run.expect(
        not cd.check_delta_consistency(
            mini_extract(
                {
                    (cd.ANCHOR_CHANNEL, cd.ANCHOR_PERIOD, "ММБ"): (100, 490),
                    (cd.ANCHOR_CHANNEL, cd.PREVIOUS_PERIOD, "ММБ"): (100, 486),
                }
            )
        ),
        "delta-consistency молчит на согласованных значениях",
    )


def check_share_digits(run: Run, extract: vs.Extract) -> None:
    """Целые проценты не воспроизводят VOC карточки — это и есть `D-40`."""
    digits = cd.SHARE_DIGITS
    try:
        cd.SHARE_DIGITS = 0
        problems = cd.check_distribution_card(extract)
    finally:
        cd.SHARE_DIGITS = digits
    run.expect(
        bool(problems),
        "distribution-card падает на целых процентах (контроль D-40)",
        str(problems[:2]),
    )


def main() -> int:
    run = Run()
    workspace = Path(tempfile.mkdtemp(prefix="voc-pipeline-"))
    sample = workspace / "voc-dashboard.xlsx"
    twin = workspace / "twin.xlsx"

    print("== генератор ==")
    first = run_tool("make_sample_data.py", "--out", str(sample))
    run_tool("make_sample_data.py", "--out", str(twin))
    run.expect(first.returncode == 0, "генератор проходит", first.stderr)
    run.expect(
        sample.read_bytes() == twin.read_bytes(),
        "два прогона дают байт-в-байт одинаковый файл",
    )
    committed = ROOT / "data" / "voc-dashboard.xlsx"
    if committed.exists():
        run.expect(
            committed.read_bytes() == sample.read_bytes(),
            "закоммиченный xlsx совпадает со свежей генерацией",
            "перегенерируйте data/voc-dashboard.xlsx",
        )

    print("\n== программа проверки ==")
    full = run_tool("check_data.py", "--source", str(sample), "--sample")
    run.expect(full.returncode == 0, "check_data проходит на корректной выгрузке", full.stdout[-600:])
    missing = [code for code in cd.ALL_CHECKS if f"      {code}\n" not in full.stdout]
    run.expect(
        not missing,
        "программа печатает каждую проверку из реестра",
        f"нет в отчёте: {missing}",
    )
    extract_only = run_tool("check_data.py", "--source", str(sample))
    leaked = [code for code, _ in cd.SAMPLE_CHECKS if f"      {code}\n" in extract_only.stdout]
    run.expect(
        not leaked,
        "без --sample проверки сэмпла не запускаются",
        f"просочились: {leaked}",
    )

    print("\n== на корректной выгрузке молчат все ==")
    base = vs.load(sample)
    for code, check in cd.ALL_CHECKS.items():
        problems = check(base)
        run.expect(not problems, f"{code} молчит", str(problems[:2])[:200])

    print("\n== каждый инвариант ломается намеренно ==")
    for code, mutation in CASES:
        broken = copy_of(base)
        mutation(broken)
        run.expect(bool(cd.ALL_CHECKS[code](broken)), f"{code} падает после «{mutation.__name__}»")
    check_arithmetic_cases(run)
    check_share_digits(run, base)

    covered = {code for code, _ in CASES} | {"delta-consistency", "distribution-card"}
    run.expect(
        not (set(cd.ALL_CHECKS) - covered),
        "у каждой проверки есть намеренная поломка",
        f"без поломки: {sorted(set(cd.ALL_CHECKS) - covered)}",
    )
    run.expect(
        not unregistered_checks(),
        "каждая написанная проверка попала в реестр",
        f"вне реестра: {unregistered_checks()}",
    )
    cd.check_never_registered = lambda extract: []
    try:
        run.expect(
            unregistered_checks() == ["check_never_registered"],
            "проверка вне реестра обнаруживается (контроль механизма)",
            f"нашлось: {unregistered_checks()}",
        )
    finally:
        del cd.check_never_registered

    print("\n== схема выгрузки ==")
    from openpyxl import load_workbook

    renamed = workspace / "renamed-column.xlsx"
    workbook = load_workbook(sample)
    workbook[cd.ANCHOR_CHANNEL].cell(row=1, column=7).value = "MARK"
    workbook.save(renamed)
    schema = run_tool("check_data.py", "--source", str(renamed))
    run.expect(
        schema.returncode != 0 and "ОШИБКА  schema" in schema.stdout,
        "переименованная колонка ловится как schema",
        schema.stdout[-300:],
    )

    text_mark = workspace / "text-mark.xlsx"
    workbook = load_workbook(sample)
    workbook[cd.ANCHOR_CHANNEL].cell(row=2, column=7).value = "5"
    workbook.save(text_mark)
    report = run_tool("check_data.py", "--source", str(text_mark))
    run.expect(
        report.returncode != 0
        and "ОШИБКА  mark-range" in report.stdout
        and "Traceback" not in report.stderr
        and "labels" in report.stdout,
        "текст в MARK1_VALUE даёт код проверки, а не traceback, и отчёт доходит до конца",
        report.stderr[-300:],
    )

    print("\n== конвертер ==")
    out = workspace / "site-data"
    convert = run_tool("xlsx_to_json.py", "--source", str(sample), "--out", str(out))
    run.expect(convert.returncode == 0, "конвертер проходит", convert.stderr)

    manifest = json.loads((out / "manifest.json").read_text())
    by_role = {file["role"]: file for file in manifest["files"]}
    ratings = json.loads((out / by_role["ratings"]["name"]).read_text())
    verbatim = json.loads((out / by_role["verbatim"]["name"]).read_text())
    reference = json.loads((out / by_role["reference"]["name"]).read_text())

    run.expect(all("answer_date" not in row for row in ratings), "ANSWER_DATE не попадает (D-06)")
    run.expect(
        any(len(row["operations"]) > 1 for row in ratings),
        "OPERATION_NAME развёрнут конвертером (§7)",
    )
    run.expect(
        any(row["problem"] is None and row["domain"] is None for row in ratings),
        "оценки без разметки сохранены как null (D-11)",
    )
    run.expect(
        verbatim["not_for_display"] == ["oslk_expertise_name"],
        "OSLK_EXPERTISE_NAME перенесён и помечен как не для показа (D-05)",
    )
    run.expect(
        isinstance(manifest["files"], list) and set(by_role) == {"ratings", "verbatim", "reference"},
        "files — массив, файл ищется по role (D-38)",
    )
    run.expect(
        all(
            file["name"].startswith(file["role"] + ".") and len(file["name"].split(".")) == 3
            for file in manifest["files"]
        ),
        "имена файлов содержат хэш содержимого (D-36)",
    )
    run.expect(
        manifest["period"]
        == {
            "from": "2026-04-01",
            "to": "2026-05-31",
            "current": cd.ANCHOR_PERIOD,
            "previous": cd.PREVIOUS_PERIOD,
        }
        and manifest["generated_at"].endswith("+00:00"),
        "manifest несёт период выгрузки, период по умолчанию и время сборки (D-35, D-41)",
        str(manifest["period"]),
    )
    run.expect(
        set(reference["labels"]) <= set(vs.LABEL_DIMENSIONS.values()),
        "измерения labels приведены к каноническим id (D-39)",
        str(list(reference["labels"])),
    )
    run.expect(
        reference["operations"]["АБМ"]["Отправка платежа"]
        != reference["operations"]["НИБ"]["Отправка платежа"],
        "справочник переносится с разбивкой по каналам (§3.2)",
    )
    run.expect(
        by_role["ratings"]["rows"] == len(ratings) == len(base.ratings),
        "число строк в манифесте совпадает с выгрузкой",
    )

    foreign = out / "site-config.json"
    foreign.write_text("{}")
    stale = out / "ratings.000000000000.json"
    stale.write_text("[]")
    run_tool("xlsx_to_json.py", "--source", str(sample), "--out", str(out))
    run.expect(not stale.exists(), "прежние версии своих файлов удаляются")
    run.expect(foreign.exists(), "чужой файл в каталоге вывода не трогается")

    broken_mark = run_tool("xlsx_to_json.py", "--source", str(text_mark), "--out", str(out))
    run.expect(
        broken_mark.returncode != 0 and "MARK1_VALUE" in broken_mark.stderr,
        "конвертер сам отказывается писать нечисловую оценку",
        broken_mark.stderr[-200:],
    )

    print("\n== несопоставленный триггер (D-12) ==")
    trimmed = workspace / "trimmed-dictionary.xlsx"
    workbook = load_workbook(sample)
    sheet = workbook[vs.DICTIONARY_SHEET_PREFIX + cd.ANCHOR_CHANNEL]
    dropped = sheet.cell(row=2, column=1).value
    sheet.delete_rows(2)
    workbook.save(trimmed)

    strict = run_tool("xlsx_to_json.py", "--source", str(trimmed), "--out", str(out))
    run.expect(
        strict.returncode != 0 and dropped in strict.stderr,
        "конвертер падает и печатает несопоставленные значения",
        strict.stderr[-300:],
    )
    relaxed = run_tool(
        "xlsx_to_json.py", "--source", str(trimmed), "--out", str(out), "--allow-unmapped"
    )
    run.expect(relaxed.returncode == 0, "--allow-unmapped не роняет сборку", relaxed.stderr)
    manifest = json.loads((out / "manifest.json").read_text())
    reference = json.loads(
        (out / next(f["name"] for f in manifest["files"] if f["role"] == "reference")).read_text()
    )
    run.expect(
        reference["operations"][cd.ANCHOR_CHANNEL][dropped]["product"] == "Без сопоставления"
        and any(item["operation"] == dropped for item in manifest["unmapped_operations"]),
        "несопоставленный триггер попадает в группу «Без сопоставления»",
    )
    checked = run_tool("check_data.py", "--source", str(trimmed))
    run.expect(
        "ОШИБКА  unmapped-trigger" in checked.stdout,
        "check_data тоже ловит отставший справочник",
        checked.stdout[-300:],
    )

    print("\n== нормализация измерений (D-39) ==")
    typo = workspace / "typo-dimension.xlsx"
    workbook = load_workbook(sample)
    sheet = workbook[vs.LABELS_SHEET]
    written = sheet.cell(row=2, column=1).value
    sheet.cell(row=2, column=1).value = f"{written}л"
    workbook.save(typo)

    broken = run_tool("xlsx_to_json.py", "--source", str(typo), "--out", str(out))
    run.expect(
        broken.returncode != 0 and f"{written}л" in broken.stderr,
        "опечатка в измерении роняет конвертер и печатает написанное значение",
        broken.stderr[-300:],
    )
    checked_labels = run_tool("check_data.py", "--source", str(typo))
    run.expect(
        checked_labels.returncode != 0,
        "check_data тоже ловит опечатку в измерении",
        (checked_labels.stdout + checked_labels.stderr)[-300:],
    )

    print("\n== порог объёма (D-37) ==")
    over = run_tool(
        "xlsx_to_json.py", "--source", str(sample), "--out", str(out), "--max-gzip-mb", "0.01"
    )
    run.expect(
        "ВНИМАНИЕ" in over.stderr and "предагрегировать" in over.stderr,
        "превышение порога gzip даёт предупреждение",
        over.stderr[-200:],
    )
    under = run_tool("xlsx_to_json.py", "--source", str(sample), "--out", str(out))
    run.expect("ВНИМАНИЕ" not in under.stderr, "в пределах порога предупреждения нет")

    print()
    if run.failed:
        print(f"провалено проверок: {run.failed}", file=sys.stderr)
        return 1
    print(f"все проверки пройдены, рабочая копия: {workspace}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
