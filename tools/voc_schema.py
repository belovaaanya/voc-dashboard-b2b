"""Схема выгрузки VOC: имена листов и колонок, чтение xlsx.

Один источник имён для генератора, проверки и конвертера: разъехавшиеся
константы в трёх скриптах проявились бы уже в браузере.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

RATING_COLUMNS = [
    "VOC_CCODE",
    "APPEAL_DATE",
    "ANSWER_DATE",
    "CHANNEL_NAME",
    "CLIENTSEGMENT_CCODE",
    "OPERATION_NAME",
    "MARK1_VALUE",
    "CLIENT_COMMENT",
    "OSLK_EXPERTISE_NAME",
    "domain",
    "problem_type",
    "problem",
]

DICTIONARY_SHEET_PREFIX = "Справочник "
DICTIONARY_COLUMNS = ["триггер", "КП", "продукт"]

PLAN_SHEET = "plan"
PLAN_COLUMNS = ["канал", "сегмент", "период с", "период по", "min", "max"]

LABELS_SHEET = "labels"
LABELS_COLUMNS = ["измерение", "код", "подпись"]

# Аналитик пишет измерение по-русски (D-02), интерфейс знает канонические id.
# Словарь живёт здесь, а не в браузере (D-39), и общий для конвертера и проверки.
LABEL_DIMENSIONS = {
    "канал": "channel",
    "сегмент": "segment",
    "триггер": "trigger",
    "кп": "cp",
    "продукт": "product",
    "область": "domain",
    "тип проблемы": "problem_type",
    "проблема": "problem",
}

OPERATION_SEPARATOR = ";"

SEGMENTS = ["ММБ", "СБ", "КИБ"]
MARKS = [1, 2, 3, 4, 5]


class SchemaError(Exception):
    """Выгрузка не соответствует docs/data-model.md §3."""


@dataclass(frozen=True)
class Rating:
    sheet: str
    row: int
    voc_ccode: str
    appeal_date: date | None
    answer_date: date | None
    channel: str
    segment: str
    operation_name: str
    mark: object
    client_comment: str | None
    oslk_expertise_name: str | None
    domain: str | None
    problem_type: str | None
    problem: str | None

    @property
    def operations(self) -> list[str]:
        return split_operations(self.operation_name)


@dataclass(frozen=True)
class PlanRow:
    row: int
    channel: str
    segment: str | None
    period_from: date | None
    period_to: date | None
    minimum: object
    maximum: object


@dataclass(frozen=True)
class LabelRow:
    row: int
    dimension: str
    code: str
    label: str | None


@dataclass
class Extract:
    ratings: list[Rating]
    dictionaries: dict[str, dict[str, tuple[str, str]]]
    plan: list[PlanRow]
    labels: list[LabelRow]
    sheet_names: list[str]

    @property
    def channels(self) -> list[str]:
        seen: list[str] = []
        for rating in self.ratings:
            if rating.channel not in seen:
                seen.append(rating.channel)
        return seen


def split_operations(value: object) -> list[str]:
    """Многозначный OPERATION_NAME → список триггеров (data-model §5.2)."""
    if value is None:
        return []
    result: list[str] = []
    for part in str(value).split(OPERATION_SEPARATOR):
        trigger = part.strip()
        if trigger and trigger not in result:
            result.append(trigger)
    return result


def as_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = as_text(value)
    if text is None:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def _header(sheet) -> list[str | None]:
    for row in sheet.iter_rows(min_row=1, max_row=1, values_only=True):
        return [as_text(cell) for cell in row]
    return []


def load(path: str | Path) -> Extract:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_names = [sheet.title for sheet in workbook.worksheets]
    ratings: list[Rating] = []
    dictionaries: dict[str, dict[str, tuple[str, str]]] = {}
    plan: list[PlanRow] = []
    labels: list[LabelRow] = []

    for sheet in workbook.worksheets:
        header = _header(sheet)
        name = sheet.title
        if name == PLAN_SHEET:
            _require_header(name, header, PLAN_COLUMNS)
            plan.extend(_read_plan(sheet))
        elif name == LABELS_SHEET:
            _require_header(name, header, LABELS_COLUMNS)
            labels.extend(_read_labels(sheet))
        elif name.startswith(DICTIONARY_SHEET_PREFIX):
            _require_header(name, header, DICTIONARY_COLUMNS)
            channel = name[len(DICTIONARY_SHEET_PREFIX):].strip()
            dictionaries[channel] = _read_dictionary(sheet)
        elif header[: len(RATING_COLUMNS)] == RATING_COLUMNS:
            ratings.extend(_read_ratings(sheet))
        else:
            raise SchemaError(
                f"лист «{name}»: не лист оценок, не «{DICTIONARY_SHEET_PREFIX}<канал>», "
                f"не «{PLAN_SHEET}», не «{LABELS_SHEET}»; колонки: {header}"
            )

    workbook.close()
    if not ratings:
        raise SchemaError("в выгрузке нет ни одного листа оценок")
    return Extract(ratings, dictionaries, plan, labels, sheet_names)


def _require_header(sheet_name: str, header: list[str | None], expected: list[str]) -> None:
    if header[: len(expected)] != expected:
        raise SchemaError(
            f"лист «{sheet_name}»: колонки {header[: len(expected)]}, ожидались {expected}"
        )


def _read_ratings(sheet):
    for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(cell is None for cell in row):
            continue
        cells = list(row) + [None] * (len(RATING_COLUMNS) - len(row))
        yield Rating(
            sheet=sheet.title,
            row=number,
            voc_ccode=as_text(cells[0]) or "",
            appeal_date=as_date(cells[1]),
            answer_date=as_date(cells[2]),
            channel=as_text(cells[3]) or "",
            segment=as_text(cells[4]) or "",
            operation_name=as_text(cells[5]) or "",
            mark=cells[6],
            client_comment=as_text(cells[7]),
            oslk_expertise_name=as_text(cells[8]),
            domain=as_text(cells[9]),
            problem_type=as_text(cells[10]),
            problem=as_text(cells[11]),
        )


def _read_dictionary(sheet) -> dict[str, tuple[str, str]]:
    mapping: dict[str, tuple[str, str]] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        trigger = as_text(row[0]) if row else None
        if trigger is None:
            continue
        mapping[trigger] = (as_text(row[1]) or "", as_text(row[2]) or "")
    return mapping


def _read_plan(sheet):
    for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(cell is None for cell in row):
            continue
        cells = list(row) + [None] * (len(PLAN_COLUMNS) - len(row))
        yield PlanRow(
            row=number,
            channel=as_text(cells[0]) or "",
            segment=as_text(cells[1]),
            period_from=as_date(cells[2]),
            period_to=as_date(cells[3]),
            minimum=cells[4],
            maximum=cells[5],
        )


def _read_labels(sheet):
    for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(cell is None for cell in row):
            continue
        cells = list(row) + [None] * (len(LABELS_COLUMNS) - len(row))
        written = as_text(cells[0]) or ""
        dimension = LABEL_DIMENSIONS.get(written.strip().lower())
        if dimension is None:
            known = ", ".join(sorted(LABEL_DIMENSIONS))
            raise SchemaError(
                f"лист {LABELS_SHEET!r}, строка {number}: неизвестное измерение "
                f"{written!r}. Известные: {known}. Опечатка здесь молча лишила бы "
                f"интерфейс подписей, поэтому разбор остановлен (D-39)"
            )
        yield LabelRow(
            row=number,
            dimension=dimension,
            code=as_text(cells[1]) or "",
            label=as_text(cells[2]),
        )
